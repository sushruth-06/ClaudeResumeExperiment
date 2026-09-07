from __future__ import annotations

from types import SimpleNamespace

from jobsearch.resume import parser


class FakeToolUseBlock:
    type = "tool_use"
    name = "record_resume"

    def __init__(self, input_data):
        self.input = input_data


class FakeClient:
    def __init__(self, tool_input):
        self.tool_input = tool_input
        self.messages = SimpleNamespace(create=self._create)

    def _create(self, **kwargs):
        return SimpleNamespace(content=[FakeToolUseBlock(self.tool_input)])


SAMPLE_RESUME_JSON = {
    "contact": {"name": "Jordan Rivera", "email": "jordan@example.com"},
    "summary": "Backend engineer with 5 years experience.",
    "skills": ["Python", "SQL", "AWS"],
    "experience": [
        {
            "company": "Acme Corp",
            "title": "Software Engineer",
            "start_date": "Jan 2021",
            "end_date": "Present",
            "bullets": ["Built a payments service handling $2M/day."],
        }
    ],
    "education": [{"school": "State University", "degree": "B.S.", "field": "CS"}],
    "projects": [],
    "certifications": [],
}


def test_parse_resume_returns_structured_resume():
    fake = FakeClient(SAMPLE_RESUME_JSON)
    resume = parser.parse_resume("Jordan Rivera\nSoftware Engineer at Acme...", client=fake)

    assert resume.contact.name == "Jordan Rivera"
    assert "Python" in resume.skills
    assert resume.experience[0].company == "Acme Corp"
    assert resume.experience[0].bullets[0].startswith("Built a payments service")


def test_extract_text_txt(tmp_path):
    p = tmp_path / "resume.txt"
    p.write_text("Hello resume")
    assert parser.extract_text(p) == "Hello resume"


def test_extract_text_docx(tmp_path):
    import docx

    p = tmp_path / "resume.docx"
    doc = docx.Document()
    doc.add_paragraph("Jordan Rivera")
    doc.add_paragraph("Software Engineer")
    doc.save(str(p))

    text = parser.extract_text(p)
    assert "Jordan Rivera" in text
    assert "Software Engineer" in text


def test_extract_text_pdf(tmp_path):
    from pypdf import PdfWriter

    p = tmp_path / "resume.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    with open(p, "wb") as f:
        writer.write(f)

    # Blank page extracts to empty string, but this proves the pdf codepath
    # runs end-to-end without raising.
    text = parser.extract_text(p)
    assert text == ""


def test_extract_text_rejects_unsupported_suffix(tmp_path):
    import pytest

    p = tmp_path / "resume.xyz"
    p.write_text("noop")
    with pytest.raises(ValueError):
        parser.extract_text(p)
